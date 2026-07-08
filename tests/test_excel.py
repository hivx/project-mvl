"""Tests for the 4-sheet Excel output matching the reference file."""
from datetime import date
from decimal import Decimal

import openpyxl

from hdtool.aggregate import aggregate
from hdtool.excel import (
    write_workbook, BK_TT80_HEADERS, DETAIL_SP_HEADERS, DS_HOADON_HEADERS,
    DS_SANPHAM_HEADERS,
)
from hdtool.models import Invoice, TaxLine, ProductLine

MIYANO = "0301234567"
PERIOD = dict(from_date=date(2026, 6, 1), to_date=date(2026, 6, 30))


def make_inv(so_hd, ky_hieu="C22TAA", tax_lines=None, product_lines=None,
             chua="1000000", thue="100000", tong="1100000", cong_ty=""):
    if tax_lines is None:
        tax_lines = [TaxLine("10%", Decimal(chua), Decimal(thue))]
    return Invoice(
        so_hd=so_hd, ky_hieu_mau_so="1", ky_hieu=ky_hieu, ngay_lap=date(2026, 6, 15),
        mst_ban="0312345678", ten_ban="CÔNG TY ABC", mst_mua=MIYANO, ten_mua="MIYANO",
        don_vi_tien="VND", tien_chua_thue=Decimal(chua), tien_thue=Decimal(thue),
        tong_thanh_toan=Decimal(tong), tax_lines=tax_lines, cong_ty=cong_ty,
        dia_chi_ban="123 Đường X", dia_chi_mua="456 Đường Y",
        product_lines=product_lines or [],
    )


def _sheet(out, name):
    return openpyxl.load_workbook(out)[name]


def test_four_template_sheets_present(tmp_path):
    res = aggregate([make_inv("123")], buyer_msts={MIYANO})
    out = tmp_path / "r.xlsx"
    write_workbook(res, out, **PERIOD)
    names = openpyxl.load_workbook(out).sheetnames
    for s in ("Hoa don Chi tiet SP", "BK mua vao TT80", "DS hoa don", "DS san pham"):
        assert s in names


def test_bk_tt80_header_and_row(tmp_path):
    res = aggregate([make_inv("123")], buyer_msts={MIYANO})
    out = tmp_path / "r.xlsx"
    write_workbook(res, out, **PERIOD)
    ws = _sheet(out, "BK mua vao TT80")
    assert [c.value for c in ws[4]] == BK_TT80_HEADERS          # header on row 4
    row5 = {BK_TT80_HEADERS[i]: c.value for i, c in enumerate(ws[5])}
    assert row5["Số"] == "123"
    assert row5["Mã số thuế người bán"] == "0312345678"
    assert row5["Thuế suất (%)"] == "10%"
    assert Decimal(str(row5["Giá trị HHDV mua vào chưa có thuế GTGT"])) == Decimal("1000000")


def test_bk_tt80_title_and_period(tmp_path):
    res = aggregate([make_inv("123")], buyer_msts={MIYANO})
    out = tmp_path / "r.xlsx"
    write_workbook(res, out, **PERIOD)
    ws = _sheet(out, "BK mua vao TT80")
    assert "MUA VÀO" in str(ws.cell(row=2, column=1).value).upper()
    assert "01/06/2026" in str(ws.cell(row=3, column=1).value)


def test_product_lines_expand_with_negative_discount_and_computed_tax(tmp_path):
    pls = [
        ProductLine("1", "F", "Cước vận chuyển", "Chuyến", Decimal(1), Decimal("17278"), Decimal("17278"), "8%"),
        ProductLine("3", "DISCOUNT", "Chiết khấu", "Chuyến", Decimal(1), Decimal("3704"), Decimal("3704"), "8%"),
        ProductLine("1", "P", "Phí nền tảng", "Chuyến", Decimal(1), Decimal("1340"), Decimal("1340"), "KCT"),
    ]
    inv = make_inv("555", ky_hieu="C26MHA", product_lines=pls,
                   tax_lines=[TaxLine("8%", Decimal("13574"), Decimal("1086")),
                              TaxLine("KCT", Decimal("1340"), Decimal("0"))],
                   chua="14914", thue="1086", tong="16000")
    res = aggregate([inv], buyer_msts={MIYANO})
    out = tmp_path / "r.xlsx"
    write_workbook(res, out, **PERIOD)
    ws = _sheet(out, "BK mua vao TT80")
    ten_c = BK_TT80_HEADERS.index("Tên hàng hóa, dịch vụ") + 1
    gia_c = BK_TT80_HEADERS.index("Giá trị HHDV mua vào chưa có thuế GTGT") + 1
    thue_c = BK_TT80_HEADERS.index("Tiền thuế GTGT") + 1
    by_name = {ws.cell(row=r, column=ten_c).value: r for r in range(5, 8)}
    assert set(by_name) == {"Cước vận chuyển", "Chiết khấu", "Phí nền tảng"}
    # discount: value & tax negative
    dr = by_name["Chiết khấu"]
    assert ws.cell(row=dr, column=gia_c).value == -3704
    assert ws.cell(row=dr, column=thue_c).value == -296      # round(3704*8%)=296
    # transportation: computed tax 17278*8% = 1382.24 -> 1382
    tr = by_name["Cước vận chuyển"]
    assert ws.cell(row=tr, column=thue_c).value == 1382
    # KCT line -> tax 0
    assert ws.cell(row=by_name["Phí nền tảng"], column=thue_c).value == 0


def test_ds_hoadon_one_row_per_invoice_with_type(tmp_path):
    res = aggregate([make_inv("1", ky_hieu="C26MHA"), make_inv("2", ky_hieu="K26TVB")],
                    buyer_msts={MIYANO})
    out = tmp_path / "r.xlsx"
    write_workbook(res, out, **PERIOD)
    ws = _sheet(out, "DS hoa don")
    assert [c.value for c in ws[3]] == DS_HOADON_HEADERS        # header on row 3
    loai = {ws.cell(row=r, column=1).value for r in (4, 5)}
    assert "HĐ từ máy tính tiền" in loai   # C26M** -> máy tính tiền
    assert "HĐ điện tử" in loai            # K26T** -> điện tử
    assert ws.cell(row=4, column=DS_HOADON_HEADERS.index("Có XML") + 1).value == "x"


def test_detail_sp_header_only_on_first_line(tmp_path):
    pls = [ProductLine("1", "A", "Line 1", "Cái", Decimal(1), Decimal("100"), Decimal("100"), "10%"),
           ProductLine("1", "B", "Line 2", "Cái", Decimal(1), Decimal("200"), Decimal("200"), "10%")]
    inv = make_inv("777", product_lines=pls, chua="300", thue="30", tong="330",
                   tax_lines=[TaxLine("10%", Decimal("300"), Decimal("30"))])
    res = aggregate([inv], buyer_msts={MIYANO})
    out = tmp_path / "r.xlsx"
    write_workbook(res, out, **PERIOD)
    ws = _sheet(out, "Hoa don Chi tiet SP")
    assert [c.value for c in ws[4]] == DETAIL_SP_HEADERS
    # invoice header only on first product line
    assert ws.cell(row=5, column=1).value == 1
    assert ws.cell(row=5, column=4).value == "777"
    assert ws.cell(row=6, column=1).value in (None, "")
    assert ws.cell(row=6, column=4).value in (None, "")
    # product name present on both lines
    ten_c = DETAIL_SP_HEADERS.index("TÊN HÀNG HÓA, DỊCH VỤ") + 1
    assert ws.cell(row=5, column=ten_c).value == "Line 1"
    assert ws.cell(row=6, column=ten_c).value == "Line 2"


def test_ds_sanpham_one_row_per_line(tmp_path):
    pls = [ProductLine("1", "A", "Line 1", "Cái", Decimal(1), Decimal("100"), Decimal("100"), "10%"),
           ProductLine("1", "B", "Line 2", "Cái", Decimal(1), Decimal("200"), Decimal("200"), "10%")]
    inv = make_inv("777", product_lines=pls, chua="300", thue="30", tong="330")
    res = aggregate([inv], buyer_msts={MIYANO})
    out = tmp_path / "r.xlsx"
    write_workbook(res, out, **PERIOD)
    ws = _sheet(out, "DS san pham")
    assert [c.value for c in ws[4]] == DS_SANPHAM_HEADERS
    ten_c = DS_SANPHAM_HEADERS.index("TÊN HÀNG HÓA, DỊCH VỤ") + 1
    assert ws.cell(row=5, column=ten_c).value == "Line 1"
    assert ws.cell(row=6, column=ten_c).value == "Line 2"
    # STT increments per line
    assert ws.cell(row=5, column=1).value == 1
    assert ws.cell(row=6, column=1).value == 2


def test_nguon_column_appended_to_sheets(tmp_path):
    inv = make_inv("1")
    inv.nguon = "PDF"
    res = aggregate([inv], buyer_msts={MIYANO})
    out = tmp_path / "r.xlsx"
    write_workbook(res, out, **PERIOD)
    for name, hr in [("BK mua vao TT80", 4), ("DS hoa don", 3),
                     ("Hoa don Chi tiet SP", 4), ("DS san pham", 4)]:
        ws = _sheet(out, name)
        hdr = [c.value for c in ws[hr]]
        assert hdr[-1] == "Nguồn"                                  # trailing column
        assert ws.cell(row=hr + 1, column=len(hdr)).value == "PDF"  # source on data row


def test_keeps_skipped_and_manual_sheets(tmp_path):
    res = aggregate([make_inv("1")], buyer_msts={MIYANO})
    out = tmp_path / "r.xlsx"
    write_workbook(res, out, skipped=[("x.xml", "cong_ty_khac", "cty khác")],
                   manual_entry=["scan.pdf"], **PERIOD)
    names = openpyxl.load_workbook(out).sheetnames
    assert "Bỏ qua" in names and "Cần nhập tay" in names
