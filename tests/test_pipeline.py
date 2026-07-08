"""End-to-end tests for the disk->Excel pipeline (no network)."""
import shutil
from datetime import date
from pathlib import Path

import openpyxl

from hdtool.pipeline import (
    find_xml_files, load_invoices, filter_by_period, build_report, write_report,
    find_manual_entry_pdfs,
)

FIXTURES = Path(__file__).parent / "fixtures"
MIYANO = "0301234567"


def _tree(tmp_path, *names):
    """Copy fixtures into a period folder and return the data dir."""
    data = tmp_path / "data"
    dest = data / "2026-06" / "abc"
    dest.mkdir(parents=True)
    for n in names:
        shutil.copy(FIXTURES / n, dest / n)
    return data


def test_find_xml_files_recursive(tmp_path):
    data = _tree(tmp_path, "inv_10pct.xml", "inv_other_buyer.xml")
    found = find_xml_files(data)
    assert len(found) == 2


def test_load_invoices_attaches_paths_and_sibling_pdf(tmp_path):
    data = _tree(tmp_path, "inv_10pct.xml")
    # add a sibling pdf with the same stem
    (data / "2026-06" / "abc" / "inv_10pct.pdf").write_bytes(b"%PDF fake")
    invoices, skipped = load_invoices(find_xml_files(data))
    assert len(invoices) == 1
    assert invoices[0].xml_path.endswith("inv_10pct.xml")
    assert invoices[0].pdf_path.endswith("inv_10pct.pdf")


def test_load_invoices_records_malformed(tmp_path):
    data = _tree(tmp_path, "inv_10pct.xml", "inv_malformed.xml")
    invoices, skipped = load_invoices(find_xml_files(data))
    assert len(invoices) == 1
    assert len(skipped) == 1
    assert "malformed" in skipped[0][0].lower()


def test_filter_by_period_uses_invoice_date():
    from hdtool.models import Invoice
    from decimal import Decimal

    def inv(d):
        return Invoice("1", "1", "K", d, "0312345678", "S", MIYANO, "M", "VND",
                       Decimal(0), Decimal(0), Decimal(0))

    inside = inv(date(2026, 6, 15))
    outside = inv(date(2026, 7, 20))
    undated = inv(None)
    keep, dropped = filter_by_period([inside, outside, undated],
                                     date(2026, 6, 1), date(2026, 6, 30))
    assert keep == [inside]
    assert outside in dropped and undated in dropped


def test_build_report_end_to_end(tmp_path):
    data = _tree(tmp_path, "inv_10pct.xml", "inv_other_buyer.xml")
    out = tmp_path / "bangke.xlsx"
    result, skipped, manual = build_report(data, date(2026, 6, 1), date(2026, 6, 30),
                                            buyer_msts={MIYANO})
    # build_report no longer writes; writing is a separate step
    assert not out.exists()
    write_report(result, out, "xlsx", skipped=skipped, manual_entry=manual)
    assert out.exists()
    # only the Miyano-buyer invoice survives the buyer filter
    assert {inv.so_hd for inv in result.invoices} == {"123"}
    wb = openpyxl.load_workbook(out)
    assert "BK mua vao TT80" in wb.sheetnames


def test_manual_entry_scoped_to_selected_company(tmp_path):
    # PDF-only invoices under different company folders; a miyano report must list
    # only miyano's orphan PDF, not medcons' or _chua-ro's.
    data = tmp_path / "data"
    (data / "miyano" / "2026-06" / "s1").mkdir(parents=True)
    (data / "miyano" / "2026-06" / "s1" / "mine.pdf").write_bytes(b"%PDF miyano")
    (data / "medcons" / "2026-06" / "s2").mkdir(parents=True)
    (data / "medcons" / "2026-06" / "s2" / "theirs.pdf").write_bytes(b"%PDF medcons")
    (data / "_chua-ro-cong-ty" / "2026-06" / "s3").mkdir(parents=True)
    (data / "_chua-ro-cong-ty" / "2026-06" / "s3" / "unknown.pdf").write_bytes(b"%PDF ?")
    companies = {"0109529507": "miyano", "0108186705": "medcons"}
    _r, _s, manual = build_report(data, date(2026, 6, 1), date(2026, 6, 30),
                                  buyer_msts=set(companies), companies=companies,
                                  only_company="miyano")
    names = {Path(m).name for m in manual}
    assert "mine.pdf" in names
    assert "theirs.pdf" not in names
    assert "unknown.pdf" not in names


def test_manual_entry_whole_tree_when_no_company(tmp_path):
    data = tmp_path / "data"
    (data / "miyano" / "2026-06" / "s1").mkdir(parents=True)
    (data / "miyano" / "2026-06" / "s1" / "mine.pdf").write_bytes(b"%PDF a")
    (data / "medcons" / "2026-06" / "s2").mkdir(parents=True)
    (data / "medcons" / "2026-06" / "s2" / "theirs.pdf").write_bytes(b"%PDF b")
    _r, _s, manual = build_report(data, date(2026, 6, 1), date(2026, 6, 30),
                                  buyer_msts={"0109529507"})
    names = {Path(m).name for m in manual}
    assert {"mine.pdf", "theirs.pdf"} <= names   # no company selected -> whole tree


def test_build_report_only_company_flags_other_company(tmp_path):
    data = _tree(tmp_path, "inv_10pct.xml", "inv_other_buyer.xml")
    companies = {"0301234567": "miyano", "0400000000": "medcons"}
    result, skipped, manual = build_report(
        data, date(2026, 6, 1), date(2026, 6, 30),
        buyer_msts=set(companies), companies=companies, only_company="medcons")
    # report holds only the selected company's buyer invoice
    assert {inv.so_hd for inv in result.invoices} == {"77"}
    # the other valid buyer (miyano) is grouped as cong_ty_khac, NOT sai_ben_mua
    assert any(grp == "cong_ty_khac" and ident.endswith("inv_10pct.xml")
               for ident, grp, _r in skipped)


def test_write_report_csv_creates_csv_and_siblings(tmp_path):
    data = _tree(tmp_path, "inv_10pct.xml")
    result, skipped, manual = build_report(data, date(2026, 6, 1), date(2026, 6, 30),
                                            buyer_msts={MIYANO})
    out = tmp_path / "bk.csv"
    write_report(result, out, "csv", skipped=skipped, manual_entry=manual)
    assert out.exists()
    assert out.with_name("bk-bo-qua.csv").exists()


def test_find_manual_entry_pdfs_flags_orphan_pdf(tmp_path):
    data = _tree(tmp_path, "inv_10pct.xml")
    dest = data / "2026-06" / "abc"
    (dest / "inv_10pct.pdf").write_bytes(b"%PDF paired")   # paired with the xml
    (dest / "scan_only.pdf").write_bytes(b"%PDF orphan")   # no xml -> manual entry
    invoices, _ = load_invoices(find_xml_files(data))
    manual = find_manual_entry_pdfs(data, invoices)
    assert [Path(m).name for m in manual] == ["scan_only.pdf"]


def test_build_report_separates_duplicate_forward_from_real_gap(tmp_path):
    # good invoice (parses) + a corrupt file with the SAME basename (harmless
    # forward duplicate) + a corrupt file with a UNIQUE name (real gap).
    data = _tree(tmp_path, "inv_10pct.xml")             # data/2026-06/abc/inv_10pct.xml (good)
    fwd = data / "2026-06" / "fwd"; fwd.mkdir()
    (fwd / "inv_10pct.xml").write_bytes(b"garbage-not-xml")     # dup -> trung_forward
    (data / "2026-06" / "abc" / "orphan_bad.xml").write_bytes(b"garbage-not-xml")  # gap
    out = tmp_path / "r.xlsx"
    result, skipped, manual = build_report(data, date(2026, 6, 1), date(2026, 6, 30),
                                           buyer_msts={MIYANO})
    write_report(result, out, "xlsx", skipped=skipped, manual_entry=manual)
    groups = {Path(ident).name: grp for ident, grp, _reason in skipped}
    assert groups["inv_10pct.xml"] == "trung_forward"   # same-named good original exists
    assert groups["orphan_bad.xml"] == "khong_doc_duoc"  # no original -> real gap
    wb = openpyxl.load_workbook(out)
    assert "Cần lấy tay" in wb.sheetnames
    # the real gap is listed on the actionable sheet; the duplicate is not
    gap_cells = [c for row in wb["Cần lấy tay"].iter_rows(values_only=True) for c in row]
    joined = " ".join(str(v) for v in gap_cells)
    assert "orphan_bad.xml" in joined
    assert "inv_10pct.xml" not in joined


def test_non_invoice_xml_not_flagged_as_gap(tmp_path):
    data = _tree(tmp_path, "inv_10pct.xml")
    # a well-formed XML that is not an invoice (tax notification)
    (data / "2026-06" / "abc" / "Thong bao thue.xml").write_bytes(
        b"<?xml version='1.0' encoding='UTF-8'?><TBao><ND>x</ND></TBao>")
    out = tmp_path / "r.xlsx"
    result, skipped, manual = build_report(data, date(2026, 6, 1), date(2026, 6, 30),
                                           buyer_msts={MIYANO})
    write_report(result, out, "xlsx", skipped=skipped, manual_entry=manual)
    groups = {Path(ident).name: grp for ident, grp, _r in skipped}
    assert groups["Thong bao thue.xml"] == "khong_phai_hoa_don"
    wb = openpyxl.load_workbook(out)
    if "Cần lấy tay" in wb.sheetnames:
        joined = " ".join(str(c) for row in wb["Cần lấy tay"].iter_rows(values_only=True) for c in row)
        assert "Thong bao thue" not in joined


def test_build_report_writes_manual_entry_sheet(tmp_path):
    data = _tree(tmp_path, "inv_10pct.xml")
    dest = data / "2026-06" / "abc"
    (dest / "inv_10pct.pdf").write_bytes(b"%PDF paired")
    (dest / "scan_only.pdf").write_bytes(b"%PDF orphan")
    out = tmp_path / "r.xlsx"
    result, skipped, manual = build_report(data, date(2026, 6, 1), date(2026, 6, 30),
                                           buyer_msts={MIYANO})
    write_report(result, out, "xlsx", skipped=skipped, manual_entry=manual)
    assert any("scan_only.pdf" in m for m in manual)
    assert all("inv_10pct.pdf" not in m for m in manual)
    wb = openpyxl.load_workbook(out)
    assert "Cần nhập tay" in wb.sheetnames
